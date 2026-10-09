"""Offline parsers preserving page, section, table, sheet, row, and cell provenance."""

from __future__ import annotations

import csv
import unicodedata
from html.parser import HTMLParser
from pathlib import Path
from typing import cast

import pymupdf
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ma_due_diligence.domain import VdrDocument
from ma_due_diligence.errors import DocumentParseError, UnsupportedDocumentError
from ma_due_diligence.vdr.models import (
    DocumentFormat,
    ElementKind,
    ParsedElement,
    ParsedVdrDocument,
)

_CLAUSE_PATTERN = __import__("re").compile(r"^(?P<number>\d+(?:\.\d+)*)[.)]?\s+(?P<body>.+)")


def normalize_source_text(text: str) -> str:
    """Normalize encoding controls and line endings without rewriting values."""

    text = unicodedata.normalize("NFC", text.replace("\r\n", "\n").replace("\r", "\n"))
    return "".join(
        character
        for character in text
        if character in {"\n", "\t"} or unicodedata.category(character) != "Cc"
    )


class VdrDocumentParser:
    def parse(self, path: Path, document: VdrDocument) -> ParsedVdrDocument:
        suffix = path.suffix.casefold()
        try:
            if suffix == ".pdf":
                return self._pdf(path, document)
            if suffix in {".txt", ".md", ".markdown"}:
                format_ = DocumentFormat.MARKDOWN if suffix != ".txt" else DocumentFormat.TEXT
                return self._text(path, document, format_)
            if suffix == ".csv":
                return self._csv(path, document)
            if suffix == ".xlsx":
                return self._xlsx(path, document)
            if suffix in {".html", ".htm"}:
                return self._html(path, document)
        except DocumentParseError:
            raise
        except Exception as error:
            raise DocumentParseError(f"unable to parse {path.name}: {error}") from error
        raise UnsupportedDocumentError(f"unsupported document format: {path.suffix or '<none>'}")

    def _pdf(self, path: Path, document: VdrDocument) -> ParsedVdrDocument:
        elements: list[ParsedElement] = []
        warnings: list[str] = []
        try:
            with pymupdf.open(path) as pdf:  # type: ignore[no-untyped-call]
                if not pdf.is_pdf:
                    raise DocumentParseError("PDF contents are invalid")
                if pdf.page_count == 0:
                    raise DocumentParseError("PDF is empty")
                if pdf.needs_pass:
                    raise DocumentParseError("encrypted PDF is not supported")
                for page_index in range(pdf.page_count):
                    page = pdf.load_page(page_index)
                    text = normalize_source_text(cast(str, page.get_text("text")))
                    if not text.strip():
                        warnings.append(
                            f"Page {page_index + 1} has no extractable text; OCR not attempted."
                        )
                        continue
                    elements.extend(
                        _structured_text_elements(text, f"p{page_index + 1}", page_index + 1)
                    )
                page_count = pdf.page_count
        except DocumentParseError:
            raise
        except (pymupdf.EmptyFileError, pymupdf.FileDataError, RuntimeError) as error:
            raise DocumentParseError(f"invalid PDF: {error}") from error
        if not elements:
            raise DocumentParseError("PDF contains no extractable text")
        return ParsedVdrDocument(
            document, DocumentFormat.PDF, tuple(elements), page_count, tuple(warnings)
        )

    def _text(
        self, path: Path, document: VdrDocument, format_: DocumentFormat
    ) -> ParsedVdrDocument:
        try:
            text = normalize_source_text(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError) as error:
            raise DocumentParseError(f"unable to read UTF-8 text: {error}") from error
        elements = _structured_text_elements(text, "text", None)
        if not elements:
            raise DocumentParseError("document is empty")
        return ParsedVdrDocument(document, format_, elements, None)

    def _csv(self, path: Path, document: VdrDocument) -> ParsedVdrDocument:
        elements: list[ParsedElement] = []
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, strict=True))
        except (OSError, UnicodeDecodeError, csv.Error) as error:
            raise DocumentParseError(f"malformed CSV: {error}") from error
        nonempty = [row for row in rows if any(cell.strip() for cell in row)]
        if not nonempty:
            raise DocumentParseError("CSV is empty")
        header = " | ".join(nonempty[0])
        for row_number, row in enumerate(nonempty, start=1):
            last_column = get_column_letter(max(1, len(row)))
            text = " | ".join(row)
            context = None if row_number == 1 else header
            elements.append(
                ParsedElement(
                    f"csv-r{row_number}",
                    ElementKind.TABLE_ROW,
                    text,
                    table_id=path.stem,
                    row_number=row_number,
                    cell_range=f"A{row_number}:{last_column}{row_number}",
                    nearby_context=context,
                )
            )
        return ParsedVdrDocument(document, DocumentFormat.CSV, tuple(elements), None)

    def _xlsx(self, path: Path, document: VdrDocument) -> ParsedVdrDocument:
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as error:
            raise DocumentParseError(f"unreadable workbook: {error}") from error
        elements: list[ParsedElement] = []
        try:
            for sheet in workbook.worksheets:
                header: str | None = None
                for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                    cells = tuple("" if value is None else str(value) for value in values)
                    if not any(value.strip() for value in cells):
                        continue
                    last_column = get_column_letter(max(1, len(cells)))
                    text = " | ".join(cells)
                    if header is None:
                        header = text
                    elements.append(
                        ParsedElement(
                            f"xlsx-{sheet.title}-r{row_number}",
                            ElementKind.SHEET_ROW,
                            text,
                            section=sheet.title,
                            table_id=f"{path.stem}:{sheet.title}",
                            sheet_name=sheet.title,
                            row_number=row_number,
                            cell_range=f"A{row_number}:{last_column}{row_number}",
                            nearby_context=None if row_number == 1 else header,
                        )
                    )
        finally:
            workbook.close()
        if not elements:
            raise DocumentParseError("workbook contains no populated cells")
        return ParsedVdrDocument(document, DocumentFormat.XLSX, tuple(elements), None)

    def _html(self, path: Path, document: VdrDocument) -> ParsedVdrDocument:
        try:
            raw = path.read_text(encoding="utf-8")
        except (OSError, UnicodeDecodeError) as error:
            raise DocumentParseError(f"unable to read HTML: {error}") from error
        parser = _StructuredHtmlParser()
        try:
            parser.feed(raw)
            parser.close()
        except Exception as error:
            raise DocumentParseError(f"malformed HTML: {error}") from error
        if not parser.elements:
            raise DocumentParseError("HTML contains no readable content")
        return ParsedVdrDocument(document, DocumentFormat.HTML, tuple(parser.elements), None)


def _structured_text_elements(
    text: str, prefix: str, page_number: int | None
) -> tuple[ParsedElement, ...]:
    blocks = [" ".join(block.split()) for block in text.split("\n\n") if block.strip()]
    if len(blocks) == 1 and "\n" in text:
        blocks = [line.strip() for line in text.splitlines() if line.strip()]
    elements: list[ParsedElement] = []
    section: str | None = None
    for index, block in enumerate(blocks):
        stripped = block.lstrip("# ").strip()
        heading = block.startswith("#") or (len(stripped) < 100 and stripped.isupper())
        clause = _CLAUSE_PATTERN.match(stripped)
        table_like = "|" in stripped or "\t" in block
        if heading:
            kind = ElementKind.HEADING
            section = stripped
        elif clause:
            kind = ElementKind.CLAUSE
        elif table_like:
            kind = ElementKind.TABLE_ROW
        else:
            kind = ElementKind.PARAGRAPH
        elements.append(
            ParsedElement(
                f"{prefix}-e{index}",
                kind,
                stripped,
                page_number,
                None if page_number is None else page_number - 1,
                section,
                clause.group("number") if clause else None,
                f"{prefix}-table" if table_like else None,
            )
        )
    return tuple(elements)


class _StructuredHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.elements: list[ParsedElement] = []
        self._tag: str | None = None
        self._buffer: list[str] = []
        self._section: str | None = None
        self._row: list[str] = []
        self._cell: list[str] | None = None
        self._row_number = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag in {"h1", "h2", "h3", "p", "li"}:
            self._tag = tag
            self._buffer = []
        elif tag == "tr":
            self._row = []
        elif tag in {"td", "th"}:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._tag is not None:
            self._buffer.append(data)
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"}:
            value = " ".join("".join(self._cell or []).split())
            self._row.append(value)
            self._cell = None
        elif tag == "tr" and self._row:
            self._row_number += 1
            text = " | ".join(self._row)
            self.elements.append(
                ParsedElement(
                    f"html-row-{self._row_number}",
                    ElementKind.TABLE_ROW,
                    text,
                    section=self._section,
                    table_id="html-table",
                    row_number=self._row_number,
                    cell_range=f"row:{self._row_number}",
                )
            )
            self._row = []
        elif tag == self._tag:
            text = " ".join("".join(self._buffer).split())
            if text:
                heading = tag.startswith("h")
                if heading:
                    self._section = text
                self.elements.append(
                    ParsedElement(
                        f"html-e{len(self.elements)}",
                        ElementKind.HEADING if heading else ElementKind.PARAGRAPH,
                        text,
                        section=self._section,
                    )
                )
            self._tag = None
            self._buffer = []
